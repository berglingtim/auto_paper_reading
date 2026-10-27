"""The published example command must build both languages outside the repo cwd."""
import subprocess
import sys
from pathlib import Path

import docx


def test_example_command_builds_two_separate_verified_docx_files(tmp_path):
    root = Path(__file__).resolve().parents[1]
    proc = subprocess.run([sys.executable, str(root / 'examples' / 'build.py'),
                           '--out', str(tmp_path / 'reports')], cwd=tmp_path,
                          capture_output=True, text=True, encoding='utf-8', errors='replace')
    assert proc.returncode == 0, proc.stdout + proc.stderr
    for language, heading in [('zh', '摘要'), ('en', 'Abstract')]:
        report, = (tmp_path / 'reports' / language).glob('*.docx')
        document = docx.Document(report)
        assert heading in [p.text for p in document.paragraphs]
        assert document.element.body.xpath('.//m:oMath')
        assert len(document.tables) == 2
        assert report.with_suffix('.math_audit.json').exists()
