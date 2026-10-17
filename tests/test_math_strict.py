import pytest
from lxml import etree
from auto_paper_reading.omml import latex_to_omml, M_NS
from auto_paper_reading.cli import _build_docx, _inject_math
import docx

NS = {'m': M_NS}

@pytest.mark.parametrize('latex,name,sub,sup', [
    (r'O(\log_{k} n)', 'log', 'k', ''),
    (r'\sin^2 x', 'sin', '', '2'),
    (r'\log_{k}^{2} n', 'log', 'k', '2'),
])
def test_function_scripts_attach_to_function_name(latex, name, sub, sup):
    el = latex_to_omml(latex)
    fn = el.xpath('.//m:fName', namespaces=NS)[0]
    assert ''.join(fn.xpath('.//m:e//m:t/text()', namespaces=NS)) == name
    assert ''.join(fn.xpath('.//m:sub//m:t/text()', namespaces=NS)) == sub
    assert ''.join(fn.xpath('.//m:sup//m:t/text()', namespaces=NS)) == sup
    assert not any(c in ''.join(el.itertext()) for c in '_^')

@pytest.mark.parametrize('text', ['正文 d_k 应排成下标', '复杂度 n^2', '$x_1$ 还有 d_k'])
def test_build_rejects_unreviewed_literal_script_marks(tmp_path, text):
    with pytest.raises(ValueError):
        _build_docx({'sections': [{'blocks': [{'type': 'text', 'text': text}]}]}, tmp_path, tmp_path/'bad.docx')

def test_injection_failure_is_not_silently_accepted(tmp_path):
    d = docx.Document(); d.add_paragraph('\ue000'+r'\unknown{x}'+'\ue001')
    p = tmp_path/'bad.docx'; d.save(p)
    with pytest.raises(ValueError):
        _inject_math(p)

def test_audit_finds_residual_in_math_and_table(tmp_path):
    from auto_paper_reading.math_audit import audit_docx
    d = docx.Document(); p = d.add_paragraph()
    p._p.append(etree.fromstring(f'<m:oMath xmlns:m="{M_NS}"><m:r><m:t>log_kn</m:t></m:r></m:oMath>'))
    d.add_table(rows=1, cols=1).cell(0,0).text='n^2'
    path=tmp_path/'bad.docx'; d.save(path)
    result=audit_docx(path)
    assert len(result['issues']) == 2

def test_literal_exception_requires_exact_text_and_reason(tmp_path):
    from auto_paper_reading.math_audit import audit_docx
    d=docx.Document(); d.add_paragraph('文件 model_weights.bin'); d.add_paragraph('d_k')
    path=tmp_path/'literal.docx'; d.save(path)
    result=audit_docx(path, [{'text':'文件 model_weights.bin', 'reason':'实际文件名'}])
    assert len(result['issues']) == 1
    assert result['reviewed_literals'][0]['reason']=='实际文件名'

def test_verify_tool_rejects_residual_and_ambiguous_versions(tmp_path):
    import subprocess, sys
    from pathlib import Path
    import fitz
    tool=Path(__file__).resolve().parents[1]/'tools'/'verify_math.py'
    d=docx.Document(); d.add_paragraph('d_k'); p=tmp_path/'report.docx'; d.save(p)
    pdf=fitz.open(); pdf.new_page().insert_text((72,72), 'd_k'); pdf.save(tmp_path/'report.pdf'); pdf.close()
    result=subprocess.run([sys.executable,str(tool),str(p)],capture_output=True)
    assert result.returncode == 1
    d.save(tmp_path/'report_v2.docx')
    result=subprocess.run([sys.executable,str(tool),str(tmp_path)],capture_output=True)
    assert result.returncode == 2

@pytest.mark.parametrize('latex', [r'x_i_j', r'x^2^3', r'_k', r'x_', r'\log_{k'])
def test_malformed_script_is_rejected(latex):
    from auto_paper_reading.omml import LatexError
    with pytest.raises(LatexError):
        latex_to_omml(latex)

def test_verify_tool_passes_current_pair_and_rejects_stale_pdf(tmp_path):
    import subprocess, sys, os, time
    from pathlib import Path
    import fitz
    tool=Path(__file__).resolve().parents[1]/'tools'/'verify_math.py'
    d=docx.Document(); d.add_paragraph()._p.append(latex_to_omml(r'\log_k n'))
    p=tmp_path/'report.docx'; d.save(p)
    pdf=fitz.open(); pdf.new_page().insert_text((72,72), 'log k n'); pdf.save(tmp_path/'report.pdf'); pdf.close()
    result=subprocess.run([sys.executable,str(tool),str(p)],capture_output=True)
    assert result.returncode == 0, result.stdout.decode('utf-8')
    past=time.time()-60; os.utime(tmp_path/'report.pdf',(past,past))
    result=subprocess.run([sys.executable,str(tool),str(p)],capture_output=True)
    assert result.returncode == 1


def test_empty_script_structure_is_rejected(tmp_path):
    from auto_paper_reading.math_audit import audit_docx
    d=docx.Document(); d.add_paragraph()._p.append(etree.fromstring(
        f'<m:oMath xmlns:m="{M_NS}"><m:sSub><m:e><m:r><m:t>x</m:t></m:r></m:e><m:sub/></m:sSub></m:oMath>'))
    p=tmp_path/'empty.docx'; d.save(p)
    assert not audit_docx(p)['passed']
