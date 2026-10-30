"""Converter behavior is tested without Office; installed backends get a real smoke test."""
import json
import subprocess
from urllib.parse import unquote, urlparse
from pathlib import Path

import docx
import fitz
import pytest

from auto_paper_reading import cli
from auto_paper_reading.parser import parse_paper
from auto_paper_reading.parser import convert


def write_pdf(path, text="Converted paper"):
    with fitz.open() as pdf:
        pdf.new_page().insert_text((72, 72), text)
        pdf.save(path)


@pytest.fixture
def converter_process(monkeypatch, tmp_path):
    """Only replace the external process; keep paths, PDF checks and parsing real."""
    executable = tmp_path / "LibreOffice" / "soffice.com"
    executable.parent.mkdir()
    executable.touch()
    monkeypatch.setenv("LIBREOFFICE_PATH", str(executable))
    commands = []

    def run(cmd, **kwargs):
        commands.append(cmd)
        assert kwargs.get("shell", False) is False
        if "--outdir" in cmd:
            target = Path(cmd[cmd.index("--outdir") + 1]) / (Path(cmd[-1]).stem + ".pdf")
        else:
            assert cmd[1] == "word2pdf"
            target = Path(cmd[cmd.index("-o") + 1])
        write_pdf(target)
        return subprocess.CompletedProcess(cmd, 0, "converted", "")

    monkeypatch.setattr(convert.subprocess, "run", run)
    return executable, commands


@pytest.mark.parametrize("extension", [".docx", ".doc"])
def test_default_conversion_uses_libreoffice_with_private_profile(tmp_path, monkeypatch, converter_process, extension):
    executable, commands = converter_process
    monkeypatch.setattr(convert, "find_wps_convert", lambda: pytest.fail("Default conversion must not invoke WPS"))
    source = tmp_path / ("论文 with spaces" + extension)
    source.write_bytes(b"Office input is read only by the external process")
    result = convert.doc_to_pdf(source, tmp_path / "out")
    assert result == tmp_path / "out" / "论文 with spaces.pdf"
    with fitz.open(result) as pdf:
        assert "Converted paper" in pdf[0].get_text()
    cmd = commands[0]
    assert cmd[0] == str(executable)
    assert "--headless" in cmd and "pdf:writer_pdf_Export" in cmd
    assert Path(cmd[-1]).suffix == source.suffix
    assert not Path(cmd[-1]).exists()  # The private working copy is cleaned up.
    profile = next(arg for arg in cmd if arg.startswith("-env:UserInstallation="))
    assert profile.startswith("-env:UserInstallation=file:")
    assert "--nolockcheck" not in cmd
    convert.doc_to_pdf(source, tmp_path / "out")
    assert profile not in commands[1]


@pytest.mark.parametrize("failure", ["exit", "missing", "invalid", "timeout"])
def test_failed_conversion_preserves_previous_pdf(tmp_path, monkeypatch, converter_process, failure):
    src = tmp_path / "report.docx"
    src.touch()
    out = tmp_path / "out"
    out.mkdir()
    old = out / "report.pdf"
    write_pdf(old, "Previous verified document")
    previous = old.read_bytes()

    def fail(cmd, **kwargs):
        if failure == "timeout":
            raise subprocess.TimeoutExpired(cmd, kwargs["timeout"])
        if failure == "invalid":
            (Path(cmd[cmd.index("--outdir") + 1]) / (Path(cmd[-1]).stem + ".pdf")).write_text("Not a PDF")
        return subprocess.CompletedProcess(cmd, 1 if failure == "exit" else 0, "", "conversion error")

    monkeypatch.setattr(convert.subprocess, "run", fail)
    with pytest.raises(RuntimeError, match="invalid PDF" if failure == "invalid" else None):
        convert.doc_to_pdf(src, out, timeout=1)
    assert old.read_bytes() == previous
    assert list(out.iterdir()) == [old]


def test_no_libreoffice_does_not_silently_fall_back_to_wps(tmp_path, monkeypatch):
    src = tmp_path / "report.docx"
    src.touch()
    monkeypatch.setattr(convert, "find_libreoffice", lambda: None, raising=False)
    monkeypatch.setattr(convert, "find_wps_convert", lambda: pytest.fail("WPS requires explicit selection"))
    with pytest.raises(RuntimeError, match="LibreOffice"):
        convert.doc_to_pdf(src, tmp_path / "out")


def test_explicit_wps_conversion_remains_available(tmp_path, monkeypatch, converter_process):
    _, commands = converter_process
    monkeypatch.setattr(convert, "find_wps_convert", lambda: Path("kwpsconvert.exe"))
    src = tmp_path / "report.docx"
    src.touch()
    result = convert.doc_to_pdf(src, tmp_path / "out", backend="wps")
    assert result.exists()
    assert commands[0][1] == "word2pdf"


def test_libreoffice_env_override_is_checked_and_not_cached(tmp_path, monkeypatch):
    executable = tmp_path / "soffice"
    executable.touch()
    monkeypatch.setenv("LIBREOFFICE_PATH", str(executable))
    assert convert.find_libreoffice() == executable
    monkeypatch.setenv("LIBREOFFICE_PATH", str(tmp_path / "missing"))
    with pytest.raises(RuntimeError, match="LIBREOFFICE_PATH"):
        convert.find_libreoffice()


@pytest.mark.parametrize("platform,command", [("win32", "soffice.com"), ("linux", "libreoffice"), ("darwin", "soffice")])
def test_libreoffice_path_discovery(tmp_path, monkeypatch, platform, command):
    executable = tmp_path / command
    executable.touch()
    monkeypatch.delenv("LIBREOFFICE_PATH", raising=False)
    monkeypatch.setattr(convert.sys, "platform", platform)
    monkeypatch.setattr(convert.shutil, "which", lambda name: str(executable) if name == command else None)
    assert convert.find_libreoffice() == executable


@pytest.mark.parametrize("extension", [".docx", ".doc"])
def test_word_input_parses_without_wps_permission(tmp_path, monkeypatch, converter_process, extension):
    source = tmp_path / ("paper" + extension)
    source.touch()
    monkeypatch.setattr(convert, "find_wps_convert", lambda: pytest.fail("Unexpected WPS call"))
    paper = parse_paper(source, tmp_path / "work")
    assert paper.format == extension[1:]
    assert paper.source_path == str(source)
    assert paper.metadata["pdf_backend"] == "libreoffice"
    assert Path(paper.metadata["converted_pdf"]).is_file()


@pytest.mark.parametrize("options", [["--pdf-backend", "wps"], ["--allow-wps"]])
def test_cli_explicit_backend_and_legacy_alias_reach_word_parser(tmp_path, monkeypatch, converter_process, options):
    monkeypatch.setattr(convert, "find_wps_convert", lambda: Path("kwpsconvert.exe"))
    source = tmp_path / "paper.doc"
    source.touch()
    assert cli.main([str(source), "--out", str(tmp_path / "out"), *options]) == 0
    paper = json.loads((tmp_path / "out/paper/paper.json").read_text(encoding="utf-8"))
    assert paper["metadata"]["pdf_backend"] == "wps"


def test_verification_uses_selected_converter_and_records_it(tmp_path, monkeypatch, converter_process):
    monkeypatch.setattr(convert, "find_wps_convert", lambda: Path("kwpsconvert.exe"))
    source = tmp_path / "report.docx"
    source.touch()
    pages = cli._verify_docx(source, tmp_path, pdf_backend="wps")
    assert len(pages) == 1 and pages[0].is_file()
    audit = json.loads((tmp_path / "verify/layout_audit.json").read_text(encoding="utf-8"))
    assert audit["pdf_backend"] == "wps"


def test_cli_returns_failure_if_requested_preview_cannot_be_rendered(tmp_path, monkeypatch):
    source = tmp_path / "paper.pdf"
    write_pdf(source)
    output = tmp_path / "out"
    assert cli.main([str(source), "--out", str(output)]) == 0
    work = output / "paper"
    (work / "content.json").write_text('{"sections":[]}', encoding="utf-8")
    monkeypatch.setattr(convert, "find_libreoffice", lambda: None, raising=False)
    monkeypatch.setattr(convert, "find_wps_convert", lambda: None)
    assert cli.main([str(source), "--out", str(output), "--stage", "docx", "--verify"]) == 1
    assert list(work.glob("*.docx"))


@pytest.mark.parametrize("backend", ["libreoffice", "wps"])
def test_installed_converter_produces_readable_pdf(tmp_path, backend):
    finder = getattr(convert, "find_libreoffice", lambda: None) if backend == "libreoffice" else convert.find_wps_convert
    if finder() is None:
        pytest.skip(f"{backend} is not installed")
    source = tmp_path / "test.docx"
    document = docx.Document()
    document.add_paragraph("Readable conversion")
    document.save(source)
    pdf = convert.doc_to_pdf(source, tmp_path / "out", backend=backend)
    with fitz.open(pdf) as result:
        assert len(result) == 1
        assert "Readable conversion" in result[0].get_text()


def test_legacy_doc_is_converted_and_parsed_with_libreoffice(tmp_path):
    executable = convert.find_libreoffice()
    if executable is None:
        pytest.skip("LibreOffice is not installed")
    source = tmp_path / "legacy.docx"
    document = docx.Document()
    document.add_paragraph("Legacy Word document")
    document.save(source)
    # Generate a real OLE .doc fixture, not a renamed ZIP/DOCX.
    subprocess.run([
        str(executable), f"-env:UserInstallation={(tmp_path / 'export-profile').as_uri()}",
        "--headless", "--convert-to", "doc:MS Word 97",
        "--outdir", str(tmp_path), str(source),
    ], check=True, capture_output=True, timeout=60)
    legacy = source.with_suffix(".doc")
    assert legacy.read_bytes()[:8] == bytes.fromhex("d0cf11e0a1b11ae1")
    paper = parse_paper(legacy, tmp_path / "parsed")
    assert paper.format == "doc" and paper.metadata["pdf_backend"] == "libreoffice"
    with fitz.open(paper.metadata["converted_pdf"]) as pdf:
        assert "Legacy Word document" in pdf[0].get_text()


def test_conversion_isolated_from_deep_user_paths(tmp_path, monkeypatch, converter_process):
    """An Office process with path limits must still receive a readable input."""
    deep = tmp_path
    while len(str(deep)) < 205:
        deep = deep / "nested paper folder"
    deep.mkdir(parents=True)
    source = deep / "论文 #%.docx"
    document = docx.Document()
    document.add_paragraph("Preserved input content")
    document.save(source)
    original = source.read_bytes()
    out = deep / "output"

    def limited_office(cmd, **kwargs):
        input_path = Path(cmd[-1])
        output_dir = Path(cmd[cmd.index("--outdir") + 1])
        profile = next(arg.split("=", 1)[1] for arg in cmd
                       if arg.startswith("-env:UserInstallation="))
        paths = [str(input_path), str(output_dir), unquote(urlparse(profile).path)]
        # Model the observed external converter failure; all file I/O stays real.
        if any(len(path) > 200 for path in paths):
            return subprocess.CompletedProcess(cmd, 0, "", "")
        assert input_path.read_bytes() == original
        text = docx.Document(input_path).paragraphs[0].text
        write_pdf(output_dir / (input_path.stem + ".pdf"), text)
        return subprocess.CompletedProcess(cmd, 0, "converted", "")

    monkeypatch.setattr(convert.subprocess, "run", limited_office)
    pdf = convert.doc_to_pdf(source, out)
    assert pdf == out / "论文 #%.pdf"
    assert source.read_bytes() == original
    assert list(out.iterdir()) == [pdf]
    with fitz.open(pdf) as result:
        assert "Preserved input content" in result[0].get_text()


def test_installed_libreoffice_handles_deep_input_and_output(tmp_path):
    if convert.find_libreoffice() is None:
        pytest.skip("LibreOffice is not installed")
    deep = tmp_path
    while len(str(deep)) < 205:
        deep = deep / "nested paper folder"
    deep.mkdir(parents=True)
    source = deep / "论文 #%.docx"
    document = docx.Document()
    document.add_paragraph("Readable deep path conversion")
    document.save(source)
    original = source.read_bytes()
    pdf = convert.doc_to_pdf(source, deep / "output")
    assert source.read_bytes() == original
    with fitz.open(pdf) as result:
        assert "Readable deep path conversion" in result[0].get_text()


def test_failed_pdf_publication_copy_preserves_previous_output(tmp_path, monkeypatch, converter_process):
    source = tmp_path / "report.docx"
    source.write_bytes(b"unchanged input")
    out = tmp_path / "out"
    out.mkdir()
    previous = out / "report.pdf"
    write_pdf(previous, "Previous verified report")
    original = previous.read_bytes()
    copyfileobj = convert.shutil.copyfileobj

    def fail_copy(source_stream, output_stream, *args, **kwargs):
        if Path(output_stream.name).parent != out:
            return copyfileobj(source_stream, output_stream, *args, **kwargs)
        output_stream.write(b"partial new PDF")
        raise OSError("Destination write failed")

    monkeypatch.setattr(convert.shutil, "copyfileobj", fail_copy)
    with pytest.raises(OSError, match="Destination write failed"):
        convert.doc_to_pdf(source, out)
    assert previous.read_bytes() == original
    assert list(out.iterdir()) == [previous]
