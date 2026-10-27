"""Language selection must reach the reader-facing DOCX, not just the README."""
import json
import re

import docx
import fitz
import pytest

from auto_paper_reading.cli import _build_docx, main
from auto_paper_reading.docx_content import normalize_content
from auto_paper_reading.math_audit import audit_docx


def test_english_report_localizes_generated_labels_and_numbering(tmp_path):
    content = {
        'language': 'en', 'title': 'Attention explained',
        'meta': {'journal': 'NeurIPS', 'year': '2017'},
        'abstract': 'A short explanation with an editable formula.',
        'symbols': [{'symbol': '$Q$', 'meaning': 'Queries'}],
        'figures': [{'id': 'missing', 'path': 'missing.png', 'caption': 'Architecture'}],
        'sections': [
            {'title': 'Method', 'blocks': [
                {'type': 'heading', 'level': 2, 'text': 'Attention'},
                {'type': 'heading', 'level': 3, 'text': 'Scaling'},
                {'type': 'equation', 'latex': r'y_i = x_i^2'},
                {'type': 'figure', 'fig_id': 'missing'},
            ]},
            {'title': 'Results', 'blocks': [
                {'type': 'heading', 'level': 2, 'text': 'Scope'},
                {'type': 'heading', 'level': 3, 'text': 'Evidence'},
            ]},
        ],
    }
    path = _build_docx(content, tmp_path, tmp_path / 'english.docx')
    d = docx.Document(path)
    text = ''.join(d.element.body.xpath('.//w:t/text()'))
    assert not re.search(r'[\u3400-\u9fff\uff00-\uffef]', text)
    headings = [p.text for p in d.paragraphs if p.style and p.style.style_id.startswith('Heading')]
    assert headings == ['Abstract', 'Notation', '1 Method', '1.1 Attention',
                        '1.1.1 Scaling', '2 Results', '2.1 Scope', '2.1.1 Evidence']
    assert [c.text for c in d.tables[0].rows[0].cells] == [
        'Symbol', 'Meaning', 'Type / shape', 'Role', 'Source / scope']
    assert d.tables[0].cell(1, 2).text == 'Not specified in the paper'
    assert '(1)' in text and 'Caption: Architecture' in text
    assert 'NeurIPS · 2017' in text
    assert d.styles.element.xpath('.//w:docDefaults/w:rPrDefault/w:rPr/w:lang/@w:val') == ['en-US']
    assert audit_docx(path)['passed']
    assert d.tables[0]._tbl.xpath('./w:tblPr/w:tblBorders/w:top/@w:sz') == ['12']
    assert all(c.vertical_alignment == 1 for r in d.tables[0].rows for c in r.cells)


def test_empty_english_notation_does_not_insert_chinese_fallback(tmp_path):
    path = _build_docx({'language': 'en', 'symbols': []}, tmp_path, tmp_path / 'empty.docx')
    d = docx.Document(path)
    assert d.paragraphs[0].text == 'Notation'
    assert not re.search(r'[\u3400-\u9fff]', ''.join(p.text for p in d.paragraphs))
    assert len(d.paragraphs) == 2 and not d.tables


@pytest.mark.parametrize('language', ['fr', '', None, [], 1])
def test_invalid_language_is_rejected_before_creating_a_report(tmp_path, language):
    with pytest.raises((ValueError, RuntimeError), match='language'):
        _build_docx({'language': language}, tmp_path, tmp_path / 'invalid.docx')
    assert not (tmp_path / 'invalid.docx').exists()


def test_normalization_defaults_to_chinese_and_rejects_unsupported_languages():
    assert normalize_content({})['language'] == 'zh'
    assert normalize_content({'language': 'en'})['language'] == 'en'
    with pytest.raises(ValueError, match='language'):
        normalize_content({'language': 'de'})


@pytest.mark.parametrize('stored_language,option,prefix,heading', [
    (None, None, '讲解_', '一 Method'),
    ('en', None, 'Explanation_', '1 Method'),
    ('zh', 'en', 'Explanation_', '1 Method'),
    ('en', 'zh', '讲解_', '一 Method'),
])
def test_cli_uses_content_language_unless_explicitly_overridden(
        tmp_path, capsys, stored_language, option, prefix, heading):
    paper = tmp_path / 'paper.pdf'
    with fitz.open() as pdf:
        pdf.new_page().insert_text((72, 72), 'Language test paper')
        pdf.save(paper)
    out = tmp_path / 'reports'
    assert main([str(paper), '--out', str(out), '--stage', 'prep']) == 0
    work = out / 'paper'
    content = {'sections': [{'title': 'Method', 'blocks': [
        {'type': 'text', 'text': 'Authored prose is preserved.'}]}]}
    if stored_language is not None:
        content['language'] = stored_language
    (work / 'content.json').write_text(json.dumps(content), encoding='utf-8')
    capsys.readouterr()
    args = [str(paper), '--out', str(out), '--stage', 'docx']
    if option:
        args += ['--language', option]
    assert main(args) == 0
    path, = work.glob('*.docx')
    assert path.name.startswith(prefix)
    d = docx.Document(path)
    assert [p.text for p in d.paragraphs] == [heading, 'Authored prose is preserved.']
    if prefix == 'Explanation_':
        assert not re.search(r'[\u3400-\u9fff]', capsys.readouterr().out)
