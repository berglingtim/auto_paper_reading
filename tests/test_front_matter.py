"""Reader-facing front matter and editable tables, through the real DOCX build."""
import docx
import pytest
from docx.oxml.ns import qn

from auto_paper_reading.cli import _build_docx, _inject_math
from auto_paper_reading.math_audit import audit_docx


def test_abstract_and_symbol_table_precede_body_without_consuming_heading_numbers(tmp_path):
    content = {
        'title': '论文讲解',
        'abstract': '背景与应用。解决序列建模问题。测试集准确率为 90%。',
        'symbols': [
            {'symbol': '$X_i$', 'meaning': '第 i 个输入',
             'shape': '$N \\times d$', 'role': '输入特征', 'source': '第 3 节'},
        ],
        'sections': [{'title': '基础知识', 'blocks': [
            {'type': 'heading', 'level': 2, 'text': '特征'},
            {'type': 'equation', 'latex': 'y=x'},
        ]}],
    }
    d = docx.Document(_build_docx(content, tmp_path, tmp_path / 'front.docx'))
    texts = [p.text for p in d.paragraphs]
    assert texts[:3] == ['论文讲解', '摘要', content['abstract']]
    assert texts[3:6] == ['符号表', '1 基础知识', '1.1 特征']
    assert texts[-1].strip() == '（1）'
    body = list(d.element.body)
    assert body[4].tag == qn('w:tbl')
    table = d.tables[0]
    assert [c.text for c in table.rows[0].cells] == ['符号', '中文含义', '类型／维度', '作用', '出处／适用范围']
    assert table.cell(1, 1).text == '第 i 个输入'
    assert table.cell(1, 0)._tc.xpath('.//m:sSub')
    assert len(table.cell(1, 2)._tc.xpath('.//m:oMath')) == 1
    assert table.rows[0]._tr.xpath('./w:trPr/w:tblHeader')
    assert audit_docx(tmp_path / 'front.docx')['passed']


def test_comparison_table_preserves_cells_newlines_and_inline_math(tmp_path):
    content = {'sections': [{'title': 'Baseline 对比', 'blocks': [
        {'type': 'table', 'headers': ['方法', '核心'], 'column_weights': [1, 3],
         'rows': [['Baseline', '第一步\n输出 $x_i^2$'], ['本文', '加入 **门控**']]},
        {'type': 'equation', 'latex': 'y=x'},
    ]}]}
    path = _build_docx(content, tmp_path, tmp_path / 'table.docx')
    d = docx.Document(path)
    assert len(d.tables) == 1
    table = d.tables[0]
    assert len(table.rows) == 3
    assert table.cell(1, 0).text == 'Baseline'
    assert len(table.cell(1, 1).paragraphs) == 2
    assert table.cell(1, 1)._tc.xpath('.//m:sSubSup')
    assert any(r.bold for r in table.cell(2, 1).paragraphs[0].runs)
    assert d.paragraphs[-1].text.strip() == '（1）'
    widths = [c.width for c in table.columns]
    assert widths[1] > widths[0] * 2
    assert sum(widths) == sum(c.width for c in table.rows[0].cells)
    assert audit_docx(path)['passed']


@pytest.mark.parametrize('block', [
    {'type': 'table', 'headers': ['方法', '核心'], 'rows': [['漏了一列']]},
    {'type': 'table', 'headers': ['方法'], 'rows': [['甲']], 'column_weights': [0]},
    {'type': 'table', 'headers': [], 'rows': []},
    {'type': 'table', 'headers': ['公式'], 'rows': [['$$x^2$$']]},
])
def test_invalid_tables_fail_instead_of_silently_losing_content(tmp_path, block):
    with pytest.raises((ValueError, RuntimeError)):
        _build_docx({'sections': [{'blocks': [block]}]}, tmp_path, tmp_path / 'bad.docx')


@pytest.mark.parametrize('extra', [
    {'abstract': {'unexpected': 'object'}},
    {'symbols': [{'symbol': '$x$'}]},
])
def test_invalid_front_matter_is_not_silently_omitted(tmp_path, extra):
    with pytest.raises((ValueError, RuntimeError)):
        _build_docx(extra, tmp_path, tmp_path / 'bad.docx')


def test_empty_symbols_explains_absence_and_legacy_input_still_builds(tmp_path):
    d = docx.Document(_build_docx({'abstract': '无公式的定性研究。', 'symbols': []},
                                tmp_path, tmp_path / 'empty.docx'))
    assert any('未使用数学符号' in p.text for p in d.paragraphs)
    assert not d.tables
    d = docx.Document(_build_docx({'sections': [{'title': '方法', 'blocks': [
        {'type': 'text', 'text': '旧版正文'}]}]}, tmp_path, tmp_path / 'legacy.docx'))
    assert [p.text for p in d.paragraphs] == ['1 方法', '旧版正文']


def test_math_in_nested_table_cells_is_injected_and_audited(tmp_path):
    d = docx.Document()
    cell = d.add_table(rows=1, cols=1).cell(0, 0)
    cell.add_table(rows=1, cols=1).cell(0, 0).text = '\ue000x_i^2\ue001'
    path = tmp_path / 'nested.docx'
    d.save(path)
    assert _inject_math(path) == 1
    report = audit_docx(path)
    assert report['passed']
    assert report['formulas'][0]['subscripts'] == 1
    assert report['formulas'][0]['superscripts'] == 1


def test_invalid_math_in_table_cell_stops_generation(tmp_path):
    with pytest.raises(ValueError):
        _build_docx({'sections': [{'blocks': [
            {'type': 'table', 'headers': ['公式'], 'rows': [[r'$\unknown{x}$']]},
        ]}]}, tmp_path, tmp_path / 'bad_math.docx')


def test_inline_math_keeps_cell_and_caption_font_sizes(tmp_path):
    content = {'sections': [{'blocks': [
        {'type': 'table', 'headers': ['符号'], 'rows': [[r'$\mathrm{softmax}$']]},
        {'type': 'figure', 'fig_id': 'missing', 'note': '说明 $x_i$'},
        {'type': 'text', 'text': '正文 $x_i$'},
    ]}]}
    d = docx.Document(_build_docx(content, tmp_path, tmp_path / 'sizes.docx'))
    cell_sizes = d.tables[0].cell(1, 0)._tc.xpath('.//m:r/w:rPr/w:sz/@w:val')
    assert cell_sizes and set(cell_sizes) == {'18'}
    note = next(p for p in d.paragraphs if p.text == '说明 ')
    assert set(note._p.xpath('.//m:r/w:rPr/w:sz/@w:val')) == {'21'}
    body = next(p for p in d.paragraphs if p.text == '正文 ')
    assert set(body._p.xpath('.//m:r/w:rPr/w:sz/@w:val')) == {'24'}
