import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH

from auto_paper_reading.cli import _build_docx


def test_body_indent_is_two_characters_but_caption_and_equation_are_not(tmp_path):
    from PIL import Image
    Image.new('RGB', (100, 100), 'white').save(tmp_path/'figure.png')
    content = {'figures': [{'id': 'f', 'path': 'figure.png', 'caption': '图注'}],
               'sections': [{'title': '方法', 'blocks': [
                   {'type': 'text', 'text': '正文'},
                   {'type': 'equation', 'latex': 'y=x'},
                   {'type': 'figure', 'fig_id': 'f', 'note': '图后说明'},
               ]}]}
    d = docx.Document(_build_docx(content, tmp_path, tmp_path/'indent.docx'))
    for text, points in [('正文', 24), ('图后说明', 21)]:
        para = next(p for p in d.paragraphs if p.text == text)
        assert para.paragraph_format.first_line_indent is not None
        assert para.paragraph_format.first_line_indent.pt == points
    for p in d.paragraphs:
        if p.text in ['1 方法', '图注：图注'] or p._p.xpath('./m:oMath'):
            assert not p.paragraph_format.first_line_indent


def test_tall_figure_is_scaled_proportionally_and_keeps_caption(tmp_path):
    from PIL import Image
    Image.new('RGB', (400, 800), 'white').save(tmp_path/'tall.png')
    content = {'figures': [{'id': 'f', 'path': 'tall.png', 'width_in': 4,
                           'height_in': 8, 'caption': '架构图'}],
               'sections': [{'blocks': [{'type': 'figure', 'fig_id': 'f', 'note': '说明'}]}]}
    d = docx.Document(_build_docx(content, tmp_path, tmp_path/'fig.docx'))
    shape = d.inline_shapes[0]
    assert shape.height.inches <= 4.51
    assert abs(shape.width / shape.height - .5) < .01
    image_p = next(p for p in d.paragraphs if p._p.xpath('.//w:drawing'))
    assert image_p.paragraph_format.keep_with_next
    caption = next(p for p in d.paragraphs if p.text == '图注：架构图')
    assert caption.paragraph_format.keep_together
    assert not caption.paragraph_format.keep_with_next


def test_manual_size_and_page_break_apply_to_next_content_only(tmp_path):
    from PIL import Image
    Image.new('RGB', (400, 800), 'white').save(tmp_path/'tall.png')
    content = {'layout': {'max_figure_height_in': 5},
               'figures': [{'id': 'f', 'path': 'tall.png', 'width_in': 4, 'height_in': 8}],
               'sections': [{'blocks': [
                   {'type': 'figure', 'fig_id': 'f', 'max_height_in': 3},
                   {'type': 'page_break'},
                   {'type': 'text', 'text': '新页第一段\n第二段'},
                   {'type': 'figure', 'fig_id': 'f', 'page_break_before': True},
               ]}]}
    d = docx.Document(_build_docx(content, tmp_path, tmp_path/'manual.docx'))
    assert abs(d.inline_shapes[0].height.inches - 3) < .02
    assert abs(d.inline_shapes[1].height.inches - 5) < .02
    first = next(p for p in d.paragraphs if p.text == '新页第一段')
    second = next(p for p in d.paragraphs if p.text == '第二段')
    assert first.paragraph_format.page_break_before
    assert not second.paragraph_format.page_break_before
    assert d.paragraphs[-1].paragraph_format.page_break_before
    assert all(p.text or p._p.xpath('.//w:drawing') for p in d.paragraphs)


def test_body_and_figure_explanation_are_justified(tmp_path):
    content = {'title': '排版测试', 'sections': [{'title': '方法', 'blocks': [
        {'type': 'text', 'text': '正文内容'},
        {'type': 'figure', 'fig_id': 'missing', 'note': '图后讲解'},
    ]}]}
    d = docx.Document(_build_docx(content, tmp_path, tmp_path/'layout.docx'))
    for text in ['正文内容', '图后讲解']:
        p = next(p for p in d.paragraphs if p.text == text)
        assert p.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert d.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER


def test_equations_are_separate_numbered_and_preserve_surrounding_text(tmp_path):
    content = {'sections': [{'title': '方法', 'blocks': [
        {'type': 'text', 'text': '前文$$x^2$$后文。\n$$\ny = 2x\n$$'},
        {'type': 'equation', 'latex': r'z=\frac{1}{2}'},
    ]}, {'title': '变换', 'blocks': [
        {'type': 'figure', 'fig_id': 'missing', 'note': '说明$$a+b$$结束'},
    ]}]}
    for name in ['first.docx', 'second.docx']:
        d = docx.Document(_build_docx(content, tmp_path, tmp_path/name))
        eqs = [p for p in d.paragraphs if p._p.xpath('./m:oMath')]
        assert len(eqs) == 4
        assert [p.text.strip() for p in eqs] == ['（1）', '（2）', '（3）', '（4）']
        assert all(len(p.paragraph_format.tab_stops) == 2 for p in eqs)
        texts = [p.text for p in d.paragraphs]
        assert all(t in texts for t in ['前文', '后文。', '说明', '结束'])
        assert '\ue000' not in d._element.xml


def test_inline_math_stays_in_sentence_and_does_not_consume_numbers(tmp_path):
    content = {'sections': [{'blocks': [
        {'type': 'text', 'text': r'本文取 $h=8$，$d_k=d_v=d_{\mathrm{model}}/h=64$。'},
        {'type': 'text', 'text': '$$y=x$$'},
        {'type': 'text', 'text': r'其中 $x_i$ 为输入。'},
    ]}]}
    d = docx.Document(_build_docx(content, tmp_path, tmp_path/'inline.docx'))
    assert len(d.paragraphs) == 3
    p = d.paragraphs[0]
    assert p.text == '本文取 ，。'
    assert len(p._p.xpath('./m:oMath')) == 2
    assert p._p.xpath('.//m:sSub')
    assert d.paragraphs[1].text.strip() == '（1）'
    assert d.paragraphs[2].text == '其中  为输入。'
    assert '\ue000' not in d._element.xml
