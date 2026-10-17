"""auto_paper_reading 命令行入口（纯本地工具，不调用任何 LLM API）。

本仓库由 AI Agent（Claude Code / Codex / Qwen Code / Kimi Code 等）按 SKILL.md
工作流驱动：Agent 负责搜索下载论文、视觉选图与撰写 content.json，
这里只负责机械的两步——解析裁剪、docx 排版与自检。

用法示例：
  # 阶段 1：解析 + 裁剪图片（.doc/.docx 输入需显式 --allow-wps）
  python run.py papers/xxx.pdf --stage prep --out output/xxx

  # 阶段 2：由 Agent 撰写的 content.json 生成 docx（--verify 用 WPS 无头渲染预览）
  python run.py papers/xxx.pdf --stage docx --out output/xxx [--verify]
"""

from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
from pathlib import Path

import fitz  # pymupdf

# Windows 控制台默认 GBK，统一按 UTF-8 输出，避免中文乱码
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass


def _safe_filename(name: str, max_len: int = 80) -> str:
    """把任意字符串净化为安全的文件名（防路径分隔符/保留字符/空名）。"""
    name = re.sub(r'[\\/:*?"<>|\x00-\x1f]+', "_", name or "")
    name = name.strip().strip(". ")
    return (name[:max_len] or "paper")

from .docx_content import normalize_content, save_content
from .models import PaperDocument
from .parser import parse_paper

ROOT = Path(__file__).resolve().parents[2]
BUILD_DOCX_JS = ROOT / "tools" / "build_docx.js"


def _inject_math(docx_path: Path) -> int:
    """把 docx 中的 LaTeX 占位 run（\\uE000...\\uE001）替换为 OMML 公式。

    公式由 WPS/Word 打开时原生渲染；全部在 python-docx 层完成，不调用任何 Office 软件。
    """
    import docx as pydocx

    from .omml import MATH_MARK_E, MATH_MARK_S, latex_to_omml

    d = pydocx.Document(str(docx_path))
    count = 0
    for para in d.paragraphs:
        for run in list(para.runs):
            text = run.text or ""
            if not (text.startswith(MATH_MARK_S) and MATH_MARK_E in text):
                continue
            latex = text[len(MATH_MARK_S) : text.rfind(MATH_MARK_E)]
            try:
                omml = latex_to_omml(latex)
            except Exception as e:
                raise ValueError(f"公式解析失败，禁止交付 {latex!r}: {e}") from e
            run._r.addprevious(omml)
            run._r.getparent().remove(run._r)
            count += 1
    if count:
        d.save(str(docx_path))
    return count


def _build_docx(content: dict, work_dir: Path, out_path: Path) -> Path:
    """调用 docx-js 脚本生成 docx，再注入 OMML 公式。

    目标文件被占用（如正被 WPS 打开）时自动换新文件名重试（讲解_xxx_v2.docx 等）。
    """
    content_path = save_content(content, work_dir / "content.json")
    attempt = 0
    while True:
        target = out_path if attempt == 0 else out_path.with_name(
            out_path.stem + f"_v{attempt + 1}" + out_path.suffix
        )
        cmd = ["node", str(BUILD_DOCX_JS), str(content_path), str(target)]
        try:
            proc = subprocess.run(
                cmd, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=180
            )
        except subprocess.TimeoutExpired:
            raise RuntimeError("docx 生成超时（>180s）")
        if proc.returncode == 0:
            out_path = target
            break
        if "EBUSY" in (proc.stdout or "") + (proc.stderr or "") and attempt < 5:
            attempt += 1
            print(f"  [warn] 输出文件被占用（可能在 WPS 中打开），改用 {target.name}")
            continue
        raise RuntimeError(f"docx 生成失败:\n{proc.stdout}\n{proc.stderr}")
    n_math = _inject_math(out_path)
    if n_math:
        print(f"  [math] 注入 {n_math} 个 OMML 公式（WPS 原生渲染）")
    from .math_audit import require_math_audit
    require_math_audit(out_path, content.get('math_review_exceptions'))
    return out_path


def _check_docx_structure(docx_path: Path) -> dict:
    """不依赖任何 Office 软件的结构化校验：能打开、有段落、图片已嵌入。"""
    import zipfile

    import docx as pydocx

    d = pydocx.Document(str(docx_path))
    paras = len(d.paragraphs)
    with zipfile.ZipFile(docx_path) as z:
        media = sorted(n for n in z.namelist() if n.startswith("word/media/"))
    tables = len(d.tables)
    return {"paragraphs": paras, "tables": tables, "media": media}


def _verify_docx(docx_path: Path, work_dir: Path) -> list[Path]:
    """（可选，需显式 --verify）用 WPS 无头转换 docx→PDF 并渲染成图片检查排版。

    仅使用 kwpsconvert.exe（wpscli word2pdf，无头、不弹窗），绝不调用 wps.exe 打开 GUI。
    """
    from .parser.convert import doc_to_pdf

    verify_dir = work_dir / "verify"
    verify_dir.mkdir(parents=True, exist_ok=True)
    pdf = doc_to_pdf(docx_path, verify_dir)
    from .layout import audit_pdf
    audit = audit_pdf(pdf)
    (verify_dir / "layout_audit.json").write_text(
        json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    flagged = [p['page'] for p in audit['pages'] if p['needs_review']]
    if flagged:
        print(f"[layout] 第 {flagged} 页有较大留白，请调整图片尺寸或分页后复核")
    # 移除上一版多余页的预览，避免误把旧页面当成本次结果。
    for old in verify_dir.glob("page_*.png"):
        if re.fullmatch(r"page_\d+\.png", old.name):
            old.unlink()
    doc = fitz.open(pdf)
    imgs = []
    for pno, page in enumerate(doc):
        pix = page.get_pixmap(matrix=fitz.Matrix(110 / 72, 110 / 72), alpha=False)
        out = verify_dir / f"page_{pno + 1:02d}.png"
        pix.save(out)
        imgs.append(out)
    doc.close()
    return imgs


def _print_summary(paper: PaperDocument) -> None:
    print(f"标题: {paper.title}")
    print(f"期刊/会议: {paper.journal or '未知'}  年份: {paper.year or '未知'}")
    print(f"摘要: {paper.abstract[:200]}")
    print(f"共提取 {len(paper.figures)} 张配图:")
    for f in paper.figures:
        print(f"  - {f.id}: 第{f.page}页  {f.width_in}x{f.height_in}in  {f.caption[:80]}")


def run(args: argparse.Namespace) -> int:
    paper_path = Path(args.paper)
    if not paper_path.exists():
        print(f"[error] 论文不存在: {paper_path}", file=sys.stderr)
        return 1

    work_dir = (Path(args.out) / paper_path.stem).resolve()
    work_dir.mkdir(parents=True, exist_ok=True)

    # ---- 阶段 1：解析 + 裁剪 ----
    if args.stage == "prep":
        print(f"[1/2] 解析论文: {paper_path}")
        paper = parse_paper(paper_path, work_dir, allow_wps=getattr(args, "allow_wps", False))
        (work_dir / "paper.json").write_text(paper.to_json(), encoding="utf-8")
        _print_summary(paper)
    else:
        paper = PaperDocument.from_dict(
            json.loads((work_dir / "paper.json").read_text(encoding="utf-8"))
        )

    if args.stage == "prep":
        print(f"[done] 解析完成，产物在 {work_dir}")
        return 0

    # ---- 阶段 2：生成 docx ----
    if args.stage == "docx":
        print("[2/2] 生成 docx")
        content_path = work_dir / "content.json"
        if not content_path.exists():
            print(f"[error] 缺少 {content_path}，请先生成讲解内容", file=sys.stderr)
            return 1
        content = normalize_content(json.loads(content_path.read_text(encoding="utf-8")), work_dir)
        out_name = f"讲解_{_safe_filename(paper.title or paper_path.stem)}.docx"
        out_docx = work_dir / out_name
        out_docx = _build_docx(content, work_dir, out_docx)
        print(f"[done] docx 已生成: {out_docx}")
        # 默认只做结构化校验（不调用任何 Office 软件）
        try:
            check = _check_docx_structure(out_docx)
            print(
                f"[check] 结构校验通过：{check['paragraphs']} 段、{check['tables']} 表、"
                f"{len(check['media'])} 张嵌入图片"
            )
        except Exception as e:
            print(f"[warn] 结构校验失败: {e}")
        # 视觉渲染预览需显式 --verify（会调用 WPS 无头转换，不弹窗）
        if getattr(args, "verify", False):
            try:
                pages = _verify_docx(out_docx, work_dir)
                print(f"[verify] 渲染 {len(pages)} 页预览图到 {work_dir / 'verify'}")
            except Exception as e:
                print(f"[warn] 渲染预览失败: {e}")
        return 0

    print(f"[error] 未知阶段: {args.stage}", file=sys.stderr)
    return 1


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="auto_paper_reading：论文 → 图文讲解 docx")
    parser.add_argument("paper", help="论文路径（.pdf/.docx/.doc）")
    parser.add_argument("--out", default="output", help="输出目录（默认 output/）")
    parser.add_argument(
        "--stage",
        choices=["prep", "docx"],
        default="prep",
        help="执行哪个阶段：prep=解析裁剪 / docx=由 content.json 生成 docx 并自检",
    )
    parser.add_argument(
        "--verify",
        action="store_true",
        help="生成 docx 后用 WPS 无头转换渲染预览图（默认只做结构化校验，不调用 Office 软件）",
    )
    parser.add_argument(
        "--allow-wps",
        action="store_true",
        help="允许调用 WPS 无头转换（kwpsconvert.exe word2pdf）解析 .doc/.docx 输入。"
        "默认关闭；注意部分安全软件可能对办公软件自动化行为误报，请自行判断。",
    )
    args = parser.parse_args(argv)
    return run(args)


if __name__ == "__main__":
    sys.exit(main())
